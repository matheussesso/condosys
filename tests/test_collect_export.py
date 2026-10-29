from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from condosys import adn
from condosys.adn import AdnResponse, AdnStore, StoredDocument
from condosys.cli import app
from condosys.models import Competence
from condosys.services.collect import collect_all
from condosys.services.condominiums import active_condominiums, sync_from_certificates
from condosys.settings import Settings
from tests.conftest import PFX_PASSWORD, THE_SUN_CNPJ, make_certificate, write_pfx
from tests.test_adn_sync import FIXTURE, FakeAdn

runner = CliRunner()


@pytest.fixture
def two_condominiums(settings: Settings, the_sun_pfx: Path, tmp_path: Path) -> Settings:
    runner.invoke(app, ["cert", "add", str(the_sun_pfx)], input=f"{PFX_PASSWORD}\n")
    key, cert = make_certificate("CONDOMINIO BLOCO B:11222333000181")
    other = write_pfx(tmp_path / "bloco-b.pfx", key, cert)
    runner.invoke(app, ["cert", "add", str(other)], input=f"{PFX_PASSWORD}\n")
    summary = sync_from_certificates(settings)
    assert len(summary.added) == 2
    return settings


def test_collect_all_records_failures_and_continues(
    two_condominiums: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = two_condominiums
    fake = FakeAdn(total=3)

    def fetch(base_url: str, loaded: Any, nsu: int, *, cnpj: str | None = None) -> AdnResponse:
        if loaded.info.cnpj != THE_SUN_CNPJ:
            raise adn.AdnError("falha simulada")
        return fake(base_url, loaded, nsu, cnpj=cnpj)

    monkeypatch.setattr("condosys.adn.fetch_dfe", fetch)
    monkeypatch.setattr(
        "condosys.services.collect.sync_documents",
        lambda *a, **k: adn.sync_documents(*a, **{**k, "fetch": fetch}),
    )
    messages: list[str] = []
    results = collect_all(
        settings,
        active_condominiums(settings),
        Competence.parse("09/2026"),
        progress=lambda i, t, c, m: messages.append(f"{i}/{t} {m}"),
    )
    by_cnpj = {r.condominium_cnpj: r for r in results}
    assert by_cnpj[THE_SUN_CNPJ].success
    assert len(by_cnpj[THE_SUN_CNPJ].invoices) == 3  # same synthetic note, three NSUs
    assert not by_cnpj["11222333000181"].success
    assert by_cnpj["11222333000181"].error == "falha simulada"
    assert any(m.startswith("1/2") for m in messages)


def test_coletar_writes_workbook(two_condominiums: Settings) -> None:
    settings = two_condominiums
    AdnStore(settings.data_dir / "adn" / THE_SUN_CNPJ).save(
        StoredDocument(1, "NFSE", "K1", FIXTURE)
    )
    result = runner.invoke(
        app, ["coletar", "-c", "09/2026", "--sem-sincronizar"], env={"COLUMNS": "200"}
    )
    assert result.exit_code == 0, result.output
    (output,) = settings.output_dir.glob("notas_2026-09_*.xlsx")

    workbook = load_workbook(output)
    assert workbook.sheetnames == [
        "Notas",
        "Resumo",
        "Erros",
        "Controle PIS",
        "Controle INSS",
        "Controle ISSQN",
        "Sobre",
    ]
    notes = list(workbook["Notas"].iter_rows(min_row=2, values_only=True))
    assert len(notes) == 1
    header = [c.value for c in workbook["Notas"][1]]
    row = dict(zip(header, notes[0], strict=True))
    assert row["Nº NFS-e"] == "123"
    assert Decimal(str(row["Valor"])) == Decimal("1000")
    assert Decimal(str(row["Ref. 4,65% (PCC)"])) == Decimal("46.5")
    assert Decimal(str(row["Ref. base INSS (÷0,11)"])) == Decimal("1000")

    pis = list(workbook["Controle PIS"].iter_rows(min_row=2, values_only=True))
    assert len(pis) == 1 and pis[0][1] == "09/2026"
    inss = list(workbook["Controle INSS"].iter_rows(min_row=2, values_only=True))
    assert len(inss) == 1 and Decimal(str(inss[0][6])) == Decimal("1000")
    iss = list(workbook["Controle ISSQN"].iter_rows(min_row=2, values_only=True))
    assert len(iss) == 1 and iss[0][9] == "ISS retido"
    summary = list(workbook["Resumo"].iter_rows(min_row=2, values_only=True))
    assert {r[0] for r in summary} == {"THE SUN FULL CONDOMINIUM RESORT", "CONDOMINIO BLOCO B"}
    assert list(workbook["Erros"].iter_rows(min_row=2, values_only=True)) == []
