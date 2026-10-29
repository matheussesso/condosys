"""Excel export of a collection run.

Sheets: "Notas" (one row per invoice), "Resumo" (per condominium), "Erros" (failed
condominiums) and three sheets laid out like the control spreadsheets of the routine
(rotina-fiscal.pdf, section 6): "Controle PIS", "Controle INSS" and "Controle ISSQN".

The control rows are *proposals* built from what the invoice declares. The employee
reviews them before posting: the system flags, the employee decides.
"""

from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from condosys.documents import format_cnpj
from condosys.models import Competence, InvoiceDetail, RunResult

MONEY_FORMAT = "#,##0.00;[Red]-#,##0.00"
DATE_FORMAT = "DD/MM/YYYY"
# Gross-value rate of PIS + COFINS + CSLL (rotina-fiscal.pdf, 6.1) — reference only.
CSRF_RATE = Decimal("0.0465")
# INSS withholding rate used to derive the base (rotina-fiscal.pdf, 6.2).
INSS_RATE = Decimal("0.11")
# IBGE code of Brasília: ISS outside it goes to the ISSQN control (rotina-fiscal.pdf, 6.3).
DF_CITY_CODE = "5300108"
LAUNCHED_BY = "condosys"

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def build_output_path(output_dir: Path, competence: Competence, now: datetime) -> Path:
    """Builds ``output/notas_<AAAA-MM>_<AAAAMMDD-HHMMSS>.xlsx`` (no ``:``, Windows-safe)."""
    return output_dir / f"notas_{competence.slug}_{now:%Y%m%d-%H%M%S}.xlsx"


def _cents(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


@dataclass(frozen=True)
class _Column:
    title: str
    width: int
    kind: str = "text"  # text | money | date


def _write_sheet(
    sheet: Worksheet, columns: Sequence[_Column], rows: Sequence[Sequence[object]]
) -> None:
    sheet.append([c.title for c in columns])
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        sheet.append(list(row))
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = column.width
        if column.kind == "text":
            continue
        number_format = MONEY_FORMAT if column.kind == "money" else DATE_FORMAT
        for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
            cell.number_format = number_format
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = sheet.dimensions


def _is_outside_df(invoice: InvoiceDetail) -> bool:
    location = invoice.service_location.split(" - ")[0].strip()
    return bool(location) and location != DF_CITY_CODE


def _invoice_rows(
    results: Sequence[RunResult],
) -> list[tuple[RunResult, InvoiceDetail]]:
    return [(r, i) for r in results if r.success for i in r.invoices]


def export_results(results: Sequence[RunResult], competence: Competence, output_path: Path) -> Path:
    """Writes the run workbook.

    Args:
        results: One result per condominium.
        competence: Collected period.
        output_path: Destination file, usually from :func:`build_output_path`.

    Returns:
        The written file.
    """
    workbook = Workbook()
    notes_sheet = workbook.active
    assert notes_sheet is not None
    notes_sheet.title = "Notas"
    pairs = _invoice_rows(results)

    _write_sheet(
        notes_sheet,
        [
            _Column("Condomínio", 35),
            _Column("CNPJ condomínio", 20),
            _Column("Nº NFS-e", 10),
            _Column("Emissão", 12, "date"),
            _Column("Competência NF", 12, "date"),
            _Column("CNPJ prestador", 20),
            _Column("Prestador", 35),
            _Column("Código serviço", 12),
            _Column("Descrição", 45),
            _Column("Local prestação", 20),
            _Column("Valor", 13, "money"),
            _Column("ISSQN", 12, "money"),
            _Column("ISS retido?", 10),
            _Column("PIS", 11, "money"),
            _Column("COFINS", 11, "money"),
            _Column("CSLL/PCC retido", 13, "money"),
            _Column("INSS", 11, "money"),
            _Column("IRRF", 11, "money"),
            _Column("Deduções", 11, "money"),
            _Column("Líquido", 13, "money"),
            _Column("Ref. 4,65% (PCC)", 13, "money"),
            _Column("Ref. base INSS (÷0,11)", 14, "money"),
            _Column("Situação", 9),
            _Column("Chave de acesso", 55),
        ],
        [
            (
                r.condominium_name,
                format_cnpj(r.condominium_cnpj),
                i.number,
                i.issue_date,
                i.competence_date,
                format_cnpj(i.provider_document),
                i.provider_name,
                i.service_code or "",
                i.service_description,
                i.service_location,
                i.amount,
                i.iss_amount,
                "SIM" if i.tax_withheld else "NÃO",
                i.pis,
                i.cofins,
                i.social_contributions_withheld,
                i.inss,
                i.irrf,
                i.deductions,
                i.net_amount,
                _cents(i.amount * CSRF_RATE),
                _cents(i.inss / INSS_RATE) if i.inss else None,
                i.status,
                i.access_key or "",
            )
            for r, i in pairs
        ],
    )

    summary_rows = []
    for r in results:
        invoices = r.invoices if r.success else []
        summary_rows.append(
            (
                r.condominium_name,
                format_cnpj(r.condominium_cnpj),
                "OK" if r.success else "ERRO",
                len(invoices),
                sum((i.amount for i in invoices), Decimal("0")),
                sum((i.iss_withheld for i in invoices), Decimal("0")),
                sum((i.social_contributions_withheld for i in invoices), Decimal("0")),
                sum((i.inss for i in invoices), Decimal("0")),
                sum((i.irrf for i in invoices), Decimal("0")),
                sum((i.net_amount for i in invoices), Decimal("0")),
            )
        )
    _write_sheet(
        workbook.create_sheet("Resumo"),
        [
            _Column("Condomínio", 35),
            _Column("CNPJ", 20),
            _Column("Situação", 10),
            _Column("Qtd. notas", 10),
            _Column("Valor total", 14, "money"),
            _Column("ISS retido", 13, "money"),
            _Column("CSLL/PCC retido", 14, "money"),
            _Column("INSS retido", 13, "money"),
            _Column("IRRF", 12, "money"),
            _Column("Líquido", 14, "money"),
        ],
        summary_rows,
    )

    _write_sheet(
        workbook.create_sheet("Erros"),
        [_Column("Condomínio", 35), _Column("CNPJ", 20), _Column("Motivo", 90)],
        [
            (r.condominium_name, format_cnpj(r.condominium_cnpj), r.error or "")
            for r in results
            if not r.success
        ],
    )

    period = competence.label
    # rotina-fiscal.pdf 6.1: Período, CNPJ Prestador, Nome Prestador, NF, Valor NF,
    # Valor PIS, Data Pagamento, Venc. DARF, Lançamento.
    _write_sheet(
        workbook.create_sheet("Controle PIS"),
        [
            _Column("Condomínio", 30),
            _Column("Período", 10),
            _Column("CNPJ Prestador", 20),
            _Column("Nome Prestador", 35),
            _Column("NF", 10),
            _Column("Valor NF", 13, "money"),
            _Column("Valor PIS", 13, "money"),
            _Column("Data Pagamento", 14, "date"),
            _Column("Venc. DARF", 12, "date"),
            _Column("Lançamento", 12),
        ],
        [
            (
                r.condominium_name,
                period,
                format_cnpj(i.provider_document),
                i.provider_name,
                i.number,
                i.amount,
                i.social_contributions_withheld,
                None,  # payment date: informed by the condominium
                None,  # due date: from the DCTFWeb
                LAUNCHED_BY,
            )
            for r, i in pairs
            if i.social_contributions_withheld > 0
        ],
    )
    # rotina-fiscal.pdf 6.2: Período, CNPJ Prestador, Nome Prestador, NF, Valor NF, BC,
    # Valor INSS, Data da emissão, Venc. DARF, Lançamento.
    _write_sheet(
        workbook.create_sheet("Controle INSS"),
        [
            _Column("Condomínio", 30),
            _Column("Período", 10),
            _Column("CNPJ Prestador", 20),
            _Column("Nome Prestador", 35),
            _Column("NF", 10),
            _Column("Valor NF", 13, "money"),
            _Column("BC", 13, "money"),
            _Column("Valor INSS", 13, "money"),
            _Column("Data da emissão", 14, "date"),
            _Column("Venc. DARF", 12, "date"),
            _Column("Lançamento", 12),
        ],
        [
            (
                r.condominium_name,
                period,
                format_cnpj(i.provider_document),
                i.provider_name,
                i.number,
                i.amount,
                _cents(i.inss / INSS_RATE),
                i.inss,
                i.issue_date,
                None,
                LAUNCHED_BY,
            )
            for r, i in pairs
            if i.inss > 0
        ],
    )
    # rotina-fiscal.pdf 6.3: ISS withheld by the taker or service outside Brasília.
    _write_sheet(
        workbook.create_sheet("Controle ISSQN"),
        [
            _Column("Condomínio", 30),
            _Column("Período", 10),
            _Column("CNPJ Prestador", 20),
            _Column("Nome Prestador", 35),
            _Column("NF", 10),
            _Column("Valor NF", 13, "money"),
            _Column("ISSQN", 12, "money"),
            _Column("Retido?", 9),
            _Column("Local prestação", 22),
            _Column("Motivo", 30),
            _Column("Lançamento", 12),
        ],
        [
            (
                r.condominium_name,
                period,
                format_cnpj(i.provider_document),
                i.provider_name,
                i.number,
                i.amount,
                i.iss_amount,
                "SIM" if i.tax_withheld else "NÃO",
                i.service_location,
                "ISS retido" if i.tax_withheld else "prestação fora do DF",
                LAUNCHED_BY,
            )
            for r, i in pairs
            if i.tax_withheld or _is_outside_df(i)
        ],
    )

    info = workbook.create_sheet("Sobre")
    info.append(["Competência", period])
    info.append(["Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")])
    info.append(["Fonte", "ADN — Ambiente de Dados Nacional da NFS-e (notas tomadas)"])
    info.append(
        [
            "Atenção",
            "Valores das abas de controle são propostas a partir da nota; conferir a "
            "aplicabilidade de cada retenção antes de lançar (rotina fiscal, seção 5).",
        ]
    )
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 100

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
