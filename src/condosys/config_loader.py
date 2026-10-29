"""Reading and validating ``config/condominios.xlsx``."""

import unicodedata
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from condosys.documents import normalize_cnpj
from condosys.models import Condominium

SHEET_NAME = "condominios"

# Spreadsheet column -> model field.
COLUMNS: dict[str, str] = {
    "ativo": "active",
    "nome": "name",
    "cnpj": "cnpj",
    "certificado": "certificate",
    "observacoes": "notes",
}
REQUIRED_COLUMNS = ("ativo", "nome", "cnpj", "certificado")
# Old spreadsheets may still have "modo_acesso": every condominium now uses its own
# certificate (ADR-011), so the column is ignored.
_FIELD_TO_COLUMN = {field: column for column, field in COLUMNS.items()}

_YES = {"SIM", "S", "TRUE", "1", "X"}
_NO = {"NAO", "N", "FALSE", "0", ""}


@dataclass(frozen=True)
class RowError:
    """A validation problem located in the spreadsheet."""

    row: int | None
    column: str | None
    message: str

    def __str__(self) -> str:
        where = []
        if self.row is not None:
            where.append(f"linha {self.row}")
        if self.column:
            where.append(f"coluna '{self.column}'")
        prefix = ", ".join(where)
        return f"{prefix}: {self.message}" if prefix else self.message


class ConfigError(Exception):
    """Raised when the condominium spreadsheet is invalid."""

    def __init__(self, errors: Sequence[RowError]) -> None:
        self.errors = list(errors)
        super().__init__("\n".join(str(error) for error in self.errors))


def _plain(text: str) -> str:
    """Uppercases and strips accents and surrounding whitespace."""
    decomposed = unicodedata.normalize("NFKD", text.strip())
    return "".join(c for c in decomposed if not unicodedata.combining(c)).upper()


def _header_key(value: Any) -> str:
    return _plain(str(value or "")).lower().replace(" ", "_")


def _parse_active(value: Any) -> bool | Any:
    if isinstance(value, bool):
        return value
    text = _plain(str(value or ""))
    if text in _YES:
        return True
    if text in _NO:
        return False
    return value  # let Pydantic report it


def _cell_value(field: str, value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if field == "cnpj":
        # Integers are kept so normalize_cnpj restores the leading zeros Excel dropped.
        return value
    if isinstance(value, int | float):
        return str(value)
    return value


def load_condominiums(config_file: Path, certs_dir: Path | None = None) -> list[Condominium]:
    """Reads and validates the condominium spreadsheet.

    Args:
        config_file: Path to ``condominios.xlsx``.
        certs_dir: When given, each referenced certificate must exist in this directory.

    Returns:
        All condominiums (active and inactive), in spreadsheet order.

    Raises:
        ConfigError: With every problem found, each pointing to its row and column.
    """
    if not config_file.is_file():
        raise ConfigError([RowError(None, None, f"arquivo não encontrado: {config_file}")])
    try:
        workbook = load_workbook(config_file, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many unrelated types for corrupt files
        message = f"não foi possível abrir {config_file}: {exc}"
        raise ConfigError([RowError(None, None, message)]) from exc

    try:
        sheet = (
            workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.worksheets[0]
        )
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise ConfigError([RowError(None, None, "planilha vazia")])

    header = [_header_key(cell) for cell in rows[0]]
    missing = [column for column in REQUIRED_COLUMNS if column not in header]
    if missing:
        raise ConfigError(
            [RowError(1, None, f"colunas obrigatórias ausentes: {', '.join(missing)}")]
        )
    positions = {column: header.index(column) for column in COLUMNS if column in header}

    errors: list[RowError] = []
    condominiums: list[Condominium] = []
    for row_number, values in enumerate(rows[1:], start=2):
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        raw = {
            COLUMNS[column]: _cell_value(
                COLUMNS[column], values[index] if index < len(values) else None
            )
            for column, index in positions.items()
        }
        raw["active"] = _parse_active(raw["active"])
        raw["row"] = row_number
        try:
            condominiums.append(Condominium.model_validate(raw, context={"certs_dir": certs_dir}))
        except ValidationError as exc:
            errors.extend(_row_errors(row_number, exc))

    errors.extend(_duplicate_errors(condominiums))
    if errors:
        raise ConfigError(errors)
    return condominiums


def _row_errors(row_number: int, exc: ValidationError) -> Iterable[RowError]:
    for error in exc.errors():
        field = str(error["loc"][0]) if error["loc"] else None
        message = error["msg"].removeprefix("Value error, ")
        if field == "active":
            message = "use SIM ou NÃO"
        yield RowError(row_number, _FIELD_TO_COLUMN.get(field or "", field), message)


def _duplicate_errors(condominiums: Sequence[Condominium]) -> Iterable[RowError]:
    seen_names: dict[str, int | None] = {}
    seen_cnpjs: dict[str, int | None] = {}
    for condominium in condominiums:
        name_key = _plain(condominium.name)
        if name_key in seen_names:
            yield RowError(
                condominium.row, "nome", f"nome repetido (já usado na linha {seen_names[name_key]})"
            )
        seen_names.setdefault(name_key, condominium.row)
        if condominium.cnpj in seen_cnpjs:
            yield RowError(
                condominium.row,
                "cnpj",
                f"CNPJ repetido (já usado na linha {seen_cnpjs[condominium.cnpj]})",
            )
        seen_cnpjs.setdefault(condominium.cnpj, condominium.row)


class CondominiumNotFoundError(LookupError):
    """Raised when a lookup by name/CNPJ matches zero or several condominiums."""

    def __init__(self, query: str, candidates: Sequence[Condominium] = ()) -> None:
        self.query = query
        self.candidates = list(candidates)
        if self.candidates:
            names = ", ".join(c.name for c in self.candidates)
            message = f"'{query}' corresponde a vários condomínios: {names}"
        else:
            message = f"nenhum condomínio corresponde a '{query}'"
        super().__init__(message)


def find_condominium(condominiums: Sequence[Condominium], query: str) -> Condominium:
    """Finds a condominium by CNPJ, exact name or unique partial name.

    Matching ignores case and accents.

    Raises:
        CondominiumNotFoundError: If nothing or more than one condominium matches.
    """
    cnpj = normalize_cnpj(query)
    for condominium in condominiums:
        if condominium.cnpj == cnpj:
            return condominium
    key = _plain(query)
    exact = [c for c in condominiums if _plain(c.name) == key]
    if len(exact) == 1:
        return exact[0]
    partial = [c for c in condominiums if key in _plain(c.name)]
    if len(partial) == 1:
        return partial[0]
    raise CondominiumNotFoundError(query, partial)


EXAMPLE_ROWS: list[tuple[str, str, str, str, str]] = [
    ("SIM", "THE SUN FULL CONDOMINIUM RESORT", "12166128000109", "the-sun.pfx", ""),
    (
        "NÃO",
        "LED (EXEMPLO)",
        "12ABC34501DE35",
        "led.pfx",
        "LED: ISSQN zerado, registrar mesmo assim",
    ),
]

_HELP_LINES = (
    ("coluna", "descrição"),
    ("ativo", "SIM ou NÃO — condomínios com NÃO são ignorados na coleta"),
    ("nome", 'nome do condomínio (usado nos comandos, ex.: condosys notas "the sun")'),
    ("cnpj", "CNPJ do condomínio, com ou sem pontuação"),
    ("certificado", "nome do arquivo .pfx do próprio condomínio, dentro da pasta certs/"),
    ("observacoes", "texto livre"),
)


def write_workbook(path: Path, rows: Iterable[Sequence[Any]]) -> Path:
    """Writes a condominium spreadsheet (header, dropdowns, instructions sheet).

    Args:
        path: Destination ``.xlsx``.
        rows: Values in the order of :data:`COLUMNS`.

    Returns:
        The written file.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = SHEET_NAME
    sheet.append(list(COLUMNS))
    for row in rows:
        sheet.append(list(row))

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for column, width in zip("ABCDE", (8, 45, 20, 40, 50), strict=True):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    # CNPJ as text so Excel keeps leading zeros.
    for row_cells in sheet.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row_cells:
            cell.number_format = "@"

    active_validation = DataValidation(type="list", formula1='"SIM,NÃO"', allow_blank=False)
    sheet.add_data_validation(active_validation)
    active_validation.add("A2:A1000")

    help_sheet = workbook.create_sheet("instrucoes")
    help_sheet.column_dimensions["A"].width = 18
    help_sheet.column_dimensions["B"].width = 90
    for line in _HELP_LINES:
        help_sheet.append(list(line))
    help_sheet["A1"].font = Font(bold=True)
    help_sheet["B1"].font = Font(bold=True)

    workbook.save(path)
    return path


def create_example_workbook(path: Path) -> Path:
    """Writes the example condominium spreadsheet."""
    return write_workbook(path, EXAMPLE_ROWS)


def save_condominiums(path: Path, condominiums: Iterable[Condominium]) -> Path:
    """Writes the condominium list back to the spreadsheet (overwrites the file)."""
    return write_workbook(
        path,
        (
            ("SIM" if c.active else "NÃO", c.name, c.cnpj, c.certificate, c.notes)
            for c in condominiums
        ),
    )
